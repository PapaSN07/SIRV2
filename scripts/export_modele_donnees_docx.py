#!/usr/bin/env python3
"""Export docs/MODELE_DONNEES_SIRV2.md → .docx pour partage métier."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "docs" / "MODELE_DONNEES_SIRV2.md"
OUT_PATH = ROOT / "docs" / "MODELE_DONNEES_SIRV2.docx"


def _shade_header_cell(cell) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), "1B3A57")
    shd.set(qn("w:val"), "clear")
    tc_pr.append(shd)


def _clean_md(s: str) -> str:
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    s = re.sub(r"\*(.+?)\*", r"\1", s)
    s = re.sub(r"`(.+?)`", r"\1", s)
    return s.strip()


def _add_runs(paragraph, text: str) -> None:
    parts = re.split(r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`)", text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = "Consolas"
            run.font.size = Pt(10)
        elif part.startswith("*") and part.endswith("*") and not part.startswith("**"):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        else:
            paragraph.add_run(part)


def _flush_table(doc: Document, rows_raw: list[str]) -> None:
    rows: list[list[str]] = []
    for line in rows_raw:
        if re.match(r"^\s*\|?[\s\-:|]+\|?\s*$", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    if not rows:
        return
    ncols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncols)
    table.style = "Table Grid"
    for ri, row in enumerate(rows):
        for ci in range(ncols):
            cell = table.cell(ri, ci)
            cell.text = _clean_md(row[ci] if ci < len(row) else "")
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.font.name = "Calibri"
                    if ri == 0:
                        run.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
            if ri == 0:
                _shade_header_cell(cell)
    doc.add_paragraph()


def convert() -> Path:
    text = MD_PATH.read_text(encoding="utf-8")

    def mermaid_note(match: re.Match[str]) -> str:
        body = match.group(1).strip()
        first = body.splitlines()[0] if body else "diagramme"
        return f"\n*[Diagramme : {first} — ouvrir le .md pour le schéma graphique]*\n"

    text = re.sub(r"```mermaid\n(.*?)```", mermaid_note, text, flags=re.DOTALL)
    text = re.sub(
        r"```[a-zA-Z]*\n(.*?)```",
        lambda m: "\n".join("    " + ln for ln in m.group(1).splitlines()),
        text,
        flags=re.DOTALL,
    )

    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2)
        section.bottom_margin = Cm(2)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)

    table_buf: list[str] = []
    for line in text.splitlines():
        if "|" in line and line.strip().startswith("|"):
            table_buf.append(line)
            continue
        if table_buf:
            _flush_table(doc, table_buf)
            table_buf = []

        if line.startswith("# "):
            p = doc.add_heading(_clean_md(line[2:]), level=0)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x57)
        elif line.startswith("## "):
            p = doc.add_heading(_clean_md(line[3:]), level=1)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x57)
        elif line.startswith("### "):
            p = doc.add_heading(_clean_md(line[4:]), level=2)
            for run in p.runs:
                run.font.color.rgb = RGBColor(0x1B, 0x3A, 0x57)
        elif line.strip() == "---" or not line.strip():
            continue
        elif line.strip().startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            _add_runs(p, line.strip()[2:])
        elif line.strip().startswith("*[") and "]*" in line:
            p = doc.add_paragraph()
            run = p.add_run(_clean_md(line.strip().strip("*")))
            run.italic = True
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
        else:
            p = doc.add_paragraph()
            _add_runs(p, line.strip())

    if table_buf:
        _flush_table(doc, table_buf)

    foot = doc.add_paragraph()
    run = foot.add_run(
        "Document SIRV2 / SENELEC — export Word depuis docs/MODELE_DONNEES_SIRV2.md"
    )
    run.italic = True
    run.font.size = Pt(8)
    run.font.color.rgb = RGBColor(0x88, 0x88, 0x88)

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    return OUT_PATH


if __name__ == "__main__":
    path = convert()
    print(path.resolve())
    print(f"size={path.stat().st_size}")
