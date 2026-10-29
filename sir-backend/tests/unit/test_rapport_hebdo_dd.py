"""Rapport hebdomadaire d'exploitation DD : semaine, classement, fichiers Word / Excel / PDF."""

import zipfile
from datetime import date, datetime

from app.repositories.rapport_hebdo_repository import LigneHebdo
from app.services import rapport_hebdo_dd_service as svc
from app.services.rapport_hebdo_dd_export import exporter_tout, tableau_repartition


def _ligne(i: int, nature: str | None, cause: str | None, *, jour: int = 0, service: str = "SDDV",
           region: str = "DAKAR", end: float = 0.5) -> LigneHebdo:
    return LigneHebdo(
        id=i, numero=i, statut="valide", date_apparition=datetime(2026, 10, 17 + jour, 8, 0),
        date_remise=datetime(2026, 10, 17 + jour, 9, 0), duree_minutes=60, service_id=1, service_code=service,
        region_libelle=region, feeder="HANN 30KV", nature=nature, cause=cause, end_mwh=end, clients=100,
        localisation="Câble blessé par tiers",
    )


def test_semaine_samedi_vendredi() -> None:
    assert svc.semaine_samedi_vendredi(date(2026, 10, 24)) == (date(2026, 10, 17), date(2026, 10, 23))
    assert svc.semaine_samedi_vendredi(date(2026, 10, 23)) == (date(2026, 10, 10), date(2026, 10, 16))
    # Modèle S39-2021 : du samedi 25 septembre au vendredi 1er octobre
    assert svc.semaine_samedi_vendredi(date(2021, 10, 2)) == (date(2021, 9, 25), date(2021, 10, 1))
    assert svc.numero_semaine(date(2021, 10, 1)) == 39


def test_classement() -> None:
    c = svc.categorie
    assert c(_ligne(1, "PERMANENT", "INCONNUE")) == svc.PERMANENT
    assert c(_ligne(1, "SEMI PERMANENT", "INCONNUE")) == svc.SEMI_PERMANENT
    assert c(_ligne(1, "FUGITIF", None)) == svc.FUGITIF
    assert c(_ligne(1, "DÉLESTAGE", "INCONNUE")) == svc.DELESTAGE_HTA
    assert c(_ligne(1, "DÉLESTAGE", "MANQUE PRODUCTION")) == svc.HTB_DELESTAGE
    assert c(_ligne(1, "PERMANENT", "INCIDENT RÉSEAU HTB")) == svc.HTB_AUTRES
    assert c(_ligne(1, "COUPURE", "TRAVAUX RÉSEAU HTA")) == svc.TRAVAUX_HTA
    assert c(_ligne(1, "COUPURE", "MANŒUVRE RÉSEAU HTA")) == svc.MANOEUVRE_HTA
    assert c(_ligne(1, "COUPURE", "CIRCONSTANCIELLE")) == svc.AUTRES_HTA
    assert c(_ligne(1, "EFFACEMENT", None)) == svc.AUTRES_HTA
    assert svc.est_exogene(_ligne(1, "PERMANENT", "CHOC ACCIDENTEL"))
    assert not svc.est_exogene(_ligne(1, "PERMANENT", "INCONNUE"))


def _donnees() -> svc.RapportHebdoData:
    lignes = [
        _ligne(1, "PERMANENT", "CHOC ACCIDENTEL", end=1.0),
        _ligne(2, "FUGITIF", "INCONNUE", jour=1, service="SDRN1", region="REGIONS", end=0.2),
        _ligne(3, "COUPURE", "TRAVAUX RÉSEAU HTA", jour=2, end=0.3),
        *[_ligne(10 + i, "DÉLESTAGE", "MANQUE PRODUCTION", jour=4) for i in range(25)],
    ]
    semaine = svc._stats(lignes, date(2026, 10, 17), date(2026, 10, 23))
    semaine.saifi, semaine.saidi_h, semaine.clients_impactes, semaine.clients_reference = 0.25, 0.4, 500, 2000
    precedente = svc._stats([_ligne(99, "PERMANENT", "INCONNUE", end=5.0)], date(2026, 10, 10), date(2026, 10, 16))
    faits = svc._faits(lignes, semaine)
    return svc.RapportHebdoData(
        semaine=semaine, precedente=precedente, services=["SDDV", "SDRN1"], cumul_end_mwh=4253.0,
        cumul_end_exogenes_mwh=310.0, cumul_nombre_exogenes=778, cumul_saifi=16.4, cumul_saidi_h=21.36,
        cumul_clients_impactes=19_000_000, cumul_clients_reference=1_200_000, objectif_end_annuel_mwh=7280.0,
        faits=faits, introduction=svc._introduction(semaine, precedente, faits),
        paragraphes_semaines=[svc._paragraphe_semaine(semaine, 140.0), svc._paragraphe_semaine(precedente, 140.0)],
        dominance=svc._dominance(semaine),
    )


def test_statistiques_et_textes() -> None:
    data = _donnees()
    s = data.semaine
    # Les 25 délestages pour manque de production sortent des tableaux HTA
    assert s.total_hta.nombre == 3
    assert round(s.total_hta.end_mwh, 3) == 1.5
    assert s.exogenes.nombre == 1
    lignes = tableau_repartition(data)
    assert lignes[0] == ["", "SDDV", "SDRN1", "TOTAL"]
    assert lignes[1] == ["Permanent", "1", "0", "1"]
    assert lignes[-1][-1] == "**1 500**"
    assert "évolution contrastée" in data.introduction and "incident réseau htb" in data.introduction.lower()
    assert any("25 interruptions côté HTA" in f for f in data.faits.majeurs)
    assert data.faits.dakar and not data.faits.regions
    assert "inférieure à la valeur cible hebdomadaire de 140 MWh" in data.paragraphes_semaines[0]
    assert svc.fmt_duree_h(1 + 5 / 60) == "1:05:00"


def test_fichiers_word_excel_pdf(tmp_path) -> None:
    chemins = exporter_tout(tmp_path, "Rapport hebdo DD S43-2026", _donnees())
    assert sorted(chemins) == ["docx", "pdf", "xlsx"]
    for chemin in chemins.values():
        assert chemin.stat().st_size > 2000
    with zipfile.ZipFile(chemins["docx"]) as z:
        noms = z.namelist()
        document = z.read("word/document.xml").decode("utf-8")
    assert "word/charts/chart1.xml" in noms and "word/media/image1.png" in noms
    assert "Rapport Hebdomadaire d&apos;Exploitation" in document or "Rapport Hebdomadaire d'Exploitation" in document
    assert chemins["pdf"].read_bytes().startswith(b"%PDF")

    from openpyxl import load_workbook

    wb = load_workbook(chemins["xlsx"])
    assert wb.sheetnames == ["Synthèse quotidienne", "Répartition par SD", "Synthèse semaine"]
