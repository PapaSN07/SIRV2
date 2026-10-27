"""Export Excel partagé — rapports CRSE (Sprint 6).

# TODO(CRSE): remplacer par les gabarits officiels SAQS.
"""

from pathlib import Path

from openpyxl import Workbook

from app.services.rapport_annuel_service import AnnuelReportData
from app.services.rapport_journalier_service import JournalierReportData
from app.services.rapport_mensuel_service import MensuelDistributionReportData
from app.services.rapport_mensuel_transport_service import MensuelTransportReportData


def _journalier_scope_label(data: JournalierReportData) -> str:
    if data.service_id is not None:
        return f"Service {data.service_id}"
    if data.region_id is not None:
        return f"Région {data.region_id}"
    return "National"


def _mensuel_scope_label(data: MensuelDistributionReportData) -> str:
    if data.region_id is not None:
        return f"Région {data.region_id}"
    return data.end.aggregate.niveau.capitalize()


def export_journalier_xlsx(path: Path, data: JournalierReportData) -> None:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Synthèse"
    summary["A1"] = data.titre_periode
    summary["A2"] = f"Périmètre : {_journalier_scope_label(data)}"
    summary["A3"] = f"Interruptions : {data.nombre_interruptions}"
    summary["A4"] = f"END cumulée (MWh) : {data.end_cumulee_mwh:.4f}"
    summary["A5"] = "Template placeholder CRSE — validation SAQS requise."

    interruptions_sheet = workbook.create_sheet("Interruptions")
    interruptions_sheet.append(
        ["N°", "Type", "Statut", "Apparition", "Remise", "Feeder", "Réseau", "END (MWh)", "Clients"]
    )
    for row in data.interruptions:
        interruptions_sheet.append(
            [
                row.numero_interruption,
                row.type_direction,
                row.statut,
                row.date_apparition.isoformat(sep=" "),
                row.date_remise.isoformat(sep=" ") if row.date_remise else "",
                row.feeder_code or "",
                row.libelle_reseau or "",
                row.end_mwh,
                row.nombre_clients,
            ]
        )

    feeders_sheet = workbook.create_sheet("Feeders")
    feeders_sheet.append(["Code", "Libellé", "Interruptions", "END (MWh)"])
    for feeder in data.feeders:
        feeders_sheet.append(
            [feeder.code, feeder.libelle, feeder.nombre_interruptions, feeder.end_mwh]
        )

    workbook.save(path)


def export_mensuel_distribution_xlsx(path: Path, data: MensuelDistributionReportData) -> None:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Synthèse"
    summary["A1"] = f"Rapport mensuel distribution — {data.mois:02d}/{data.annee}"
    summary["A2"] = f"Périmètre : {_mensuel_scope_label(data)}"
    summary["A3"] = f"END (MWh) : {data.end.aggregate.end_mwh:.4f}"
    summary["A4"] = f"SAIFI : {data.saifi.aggregate.saifi:.4f}"
    summary["A5"] = f"SAIDI : {data.saidi.aggregate.saidi:.4f}"
    summary["A6"] = f"Interruptions : {data.end.aggregate.nombre_interruptions}"
    summary["A7"] = "Template placeholder CRSE — validation SAQS requise."

    services_sheet = workbook.create_sheet("Services")
    services_sheet.append(
        ["Service ID", "END (MWh)", "Interruptions", "SAIFI", "SAIDI"]
    )

    saifi_by_service = {
        item.service_id: item.saifi
        for item in (data.saifi.services.items if data.saifi.services else [])
    }
    saidi_by_service = {
        item.service_id: item.saidi
        for item in (data.saidi.services.items if data.saidi.services else [])
    }

    if data.end.services:
        for item in data.end.services.items:
            services_sheet.append(
                [
                    item.service_id,
                    item.end_mwh,
                    item.nombre_interruptions,
                    saifi_by_service.get(item.service_id),
                    saidi_by_service.get(item.service_id),
                ]
            )

    feeders_sheet = workbook.create_sheet("Top feeders END")
    feeders_sheet.append(["Rang", "Code", "Libellé", "Service", "END (MWh)", "Interruptions"])
    for rank, feeder in enumerate(data.top_feeders.items, start=1):
        feeders_sheet.append(
            [
                rank,
                feeder.code,
                feeder.libelle,
                feeder.service_id,
                feeder.end_mwh,
                feeder.nombre_interruptions,
            ]
        )

    workbook.save(path)


def export_mensuel_transport_xlsx(path: Path, data: MensuelTransportReportData) -> None:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Synthèse"
    summary["A1"] = f"Rapport mensuel transport HTB — {data.mois:02d}/{data.annee}"
    summary["A2"] = "Périmètre : National HTB"
    summary["A3"] = f"Interruptions : {data.nombre_interruptions}"
    summary["A4"] = f"END cumulée (MWh) : {data.end_cumulee_mwh:.4f}"
    summary["A5"] = "Template placeholder CRSE — validation SAQS requise."

    interruptions_sheet = workbook.create_sheet("Interruptions")
    interruptions_sheet.append(
        [
            "N°",
            "Type",
            "Statut",
            "Apparition",
            "Remise",
            "Équipement",
            "Source",
            "END (MWh)",
            "Puissance transit (MW)",
            "Durée (min)",
        ]
    )
    for row in data.interruptions:
        interruptions_sheet.append(
            [
                row.numero_interruption,
                row.type_direction,
                row.statut,
                row.date_apparition.isoformat(sep=" "),
                row.date_remise.isoformat(sep=" ") if row.date_remise else "",
                row.equipement_libelle or "",
                row.source_libelle or "",
                row.end_mwh,
                row.puissance_transit_mw,
                row.duree_minutes,
            ]
        )

    equipements_sheet = workbook.create_sheet("Top équipements END")
    equipements_sheet.append(["Rang", "Code", "Libellé", "Interruptions", "END (MWh)"])
    for rank, eq in enumerate(data.top_equipements, start=1):
        equipements_sheet.append(
            [rank, eq.code, eq.libelle, eq.nombre_interruptions, eq.end_mwh]
        )

    workbook.save(path)


def export_annuel_xlsx(path: Path, data: AnnuelReportData) -> None:
    workbook = Workbook()

    summary = workbook.active
    summary.title = "Synthèse"
    summary["A1"] = f"Rapport annuel {data.annee} — {data.direction}"
    summary["A2"] = f"END totale : {data.total_end_annee:.4f} MWh"
    summary["A3"] = data.objectifs_note
    summary["A4"] = "Template placeholder CRSE — validation SAQS requise."

    serie_sheet = workbook.create_sheet("Série mensuelle")
    serie_sheet.append(["Mois", "END (MWh)", "SAIFI", "SAIDI"])
    for point in data.serie_civile:
        serie_sheet.append([point.mois, point.end_mwh, point.saifi, point.saidi])

    if data.transport_serie:
        transport_sheet = workbook.create_sheet("Transport mensuel")
        transport_sheet.append(["Mois", "END (MWh)"])
        for point in data.transport_serie:
            transport_sheet.append([point.mois, point.end_mwh])

    if data.top_feeders:
        feeders_sheet = workbook.create_sheet("Top feeders annuel")
        feeders_sheet.append(["Rang", "Code", "Libellé", "END (MWh)", "Interruptions"])
        for rank, feeder in enumerate(data.top_feeders, start=1):
            feeders_sheet.append(
                [rank, feeder.code, feeder.libelle, feeder.end_mwh, feeder.nombre_interruptions]
            )

    if data.top_equipements:
        eq_sheet = workbook.create_sheet("Top équipements annuel")
        eq_sheet.append(["Rang", "Code", "Libellé", "END (MWh)", "Interruptions"])
        for rank, eq in enumerate(data.top_equipements, start=1):
            eq_sheet.append(
                [rank, eq.code, eq.libelle, eq.end_mwh, eq.nombre_interruptions]
            )

    if data.historique_end:
        hist_sheet = workbook.create_sheet("Historique glissant END")
        hist_sheet.append(["Année", "Mois", "Valeur"])
        for point in data.historique_end.points:
            hist_sheet.append([point.annee, point.mois, point.valeur])

    workbook.save(path)


def write_placeholder_xlsx(path: Path, *, title: str, subtitle: str) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Rapport"
    sheet["A1"] = title
    sheet["A2"] = subtitle
    sheet["A3"] = "Template placeholder CRSE — validation SAQS requise."
    workbook.save(path)
