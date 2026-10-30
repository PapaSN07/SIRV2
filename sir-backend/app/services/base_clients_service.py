"""Service — base clients SAIFI/SAIDI (reprise de l'atelier SAIFI/SAIDI de l'application bureau SIR).

Chaque mois, SIR v1 chargeait le fichier « nombredeclientsh2mc_MMAAAA.xlsx » (contrats par unité de
relève, UR) puis recalculait les clients de chaque poste (somme des contrats de ses UR) et le nombre
de clients par service (dénominateur SAIFI/SAIDI). SIRV2 lit la même donnée dans
BI_ODS.dbo.CONTRATS_PAR_AGENCE_UR (chargée par SSIS), ou dans le fichier déposé en secours.
"""

from __future__ import annotations

import asyncio
import logging
import re
from collections import defaultdict
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from sqlalchemy.ext.asyncio import AsyncSession

from app.core.config import settings
from app.repositories import clients_repository as repo

logger = logging.getLogger(__name__)

SOURCE_BI = "bi_ods"
SOURCE_FICHIER = "fichier"
_RE_UR = re.compile(r"^[0-9A-Z]{4,20}$")


class BaseClientsErreur(ValueError):
    """Fichier ou source illisible."""


@dataclass
class BilanImport:
    annee: int
    mois: int
    source: str
    unites_releve: int
    contrats: int
    contrats_sans_agence: int
    postes_mis_a_jour: int
    clients_postes: int
    services: int
    lignes_rejetees: int


# ── Lecture des sources ──────────────────────────────────────────────────────


def normaliser(brute: dict[str, Any]) -> dict[str, Any] | None:
    """Ligne commune aux deux sources ; None si l'UR ou le nombre est inexploitable."""
    ur = str(brute.get("unite_releve") or "").strip().upper()
    if not _RE_UR.match(ur):
        return None
    try:
        nb = int(float(brute.get("nb_contrats") or 0))
    except (TypeError, ValueError):
        return None
    if nb < 0:
        return None
    code_agence = str(brute.get("code_agence") or "").strip() or ur[:3]
    return {
        "unite_releve": ur,
        "libelle": (str(brute.get("libelle") or "").strip() or None),
        "code_agence": code_agence[:10],
        "delegation": (str(brute.get("delegation") or "").strip()[:10] or None),
        "nb_contrats": nb,
        "date_maj": brute.get("date_maj"),
    }


def _cle(texte: Any) -> str:
    return re.sub(r"[^a-z]", "", str(texte or "").lower())


# En-têtes du fichier H2MC : « Delegation (code) », « Agence (code) », « Unite de Releve »,
# « Unite Releve (Code) », « Nombre de contrat »
_COLONNES_FICHIER = {
    "delegationcode": "delegation",
    "agencecode": "code_agence",
    "unitedereleve": "libelle",
    "unitrelevecode": "unite_releve",
    "uniterelevecode": "unite_releve",
    "nombredecontrat": "nb_contrats",
    "nombredecontrats": "nb_contrats",
}


def lire_fichier(chemin: Path) -> list[dict[str, Any]]:
    """Lit le fichier mensuel H2MC (.xlsx) ; lève BaseClientsErreur si les colonnes manquent."""
    from openpyxl import load_workbook

    try:
        classeur = load_workbook(chemin, read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise BaseClientsErreur("Fichier Excel illisible") from exc
    feuille = classeur.worksheets[0]
    lignes = feuille.iter_rows(values_only=True)
    index: dict[str, int] = {}
    for _ in range(10):  # l'en-tête est dans les premières lignes
        entete = next(lignes, None)
        if entete is None:
            break
        index = {_COLONNES_FICHIER[_cle(v)]: i for i, v in enumerate(entete) if _cle(v) in _COLONNES_FICHIER}
        if {"unite_releve", "nb_contrats"} <= set(index):
            break
    if not {"unite_releve", "nb_contrats"} <= set(index):
        raise BaseClientsErreur("Colonnes « Unite Releve (Code) » et « Nombre de contrat » introuvables")
    brutes = [
        {cle: (row[i] if i < len(row) else None) for cle, i in index.items()} for row in lignes if row and any(row)
    ]
    classeur.close()
    return brutes


def _lire_bi_ods_sync() -> list[dict[str, Any]]:
    import pyodbc

    with pyodbc.connect(settings.bi_ods_odbc, timeout=60) as conn:
        cur = conn.cursor()
        cur.execute(
            "SELECT DELEGATION, CODE_AGENCE, AGENCE, UNITE_RELEVE, NB_CONTRATS, DATE_MAJ "
            "FROM dbo.CONTRATS_PAR_AGENCE_UR"
        )
        return [
            {
                "delegation": d,
                "code_agence": c,
                "libelle": a,
                "unite_releve": u,
                "nb_contrats": n,
                "date_maj": m,
            }
            for d, c, a, u, n, m in cur.fetchall()
        ]


async def lire_bi_ods() -> list[dict[str, Any]]:
    if not settings.bi_ods_odbc:
        raise BaseClientsErreur("Connexion BI_ODS non configurée (BI_ODS_ODBC)")
    return await asyncio.to_thread(_lire_bi_ods_sync)


# ── Import et recalcul ───────────────────────────────────────────────────────


def preparer_lignes(
    annee: int, mois: int, brutes: list[dict[str, Any]], agences: dict[str, int], source: str
) -> tuple[list[dict[str, Any]], int]:
    """Lignes prêtes à enregistrer (une par UR, doublons additionnés) et nombre de lignes rejetées."""
    par_ur: dict[str, dict[str, Any]] = {}
    rejets = 0
    for brute in brutes:
        ligne = normaliser(brute)
        if ligne is None:
            rejets += 1
            continue
        ur = ligne["unite_releve"]
        if ur in par_ur:
            par_ur[ur]["nb_contrats"] += ligne["nb_contrats"]
            continue
        par_ur[ur] = {
            **ligne,
            "annee": annee,
            "mois": mois,
            "agence_id": agences.get(ligne["code_agence"]),
            "source": source,
        }
    return list(par_ur.values()), rejets


async def importer(
    db: AsyncSession, annee: int, mois: int, brutes: list[dict[str, Any]], source: str
) -> BilanImport:
    if not (1 <= mois <= 12 and 2000 <= annee <= 2100):
        raise BaseClientsErreur("Période invalide")
    lignes, rejets = preparer_lignes(annee, mois, brutes, await repo.agences_par_code(db), source)
    if not lignes:
        raise BaseClientsErreur("Aucune unité de relève exploitable")
    await repo.remplacer_contrats(db, annee, mois, lignes)
    bilan = await recalculer(db, annee, mois)
    bilan.source = source
    bilan.unites_releve = len(lignes)
    bilan.lignes_rejetees = rejets
    await db.commit()
    logger.info("Base clients %02d/%d : %s", mois, annee, bilan)
    return bilan


async def recalculer(db: AsyncSession, annee: int, mois: int) -> BilanImport:
    """Clients de chaque poste et nombre de clients par service pour le mois (sans commit)."""
    par_agence = await repo.contrats_par_agence(db, annee, mois)
    services = await repo.services_par_agence(db)
    par_service: dict[int, int] = defaultdict(int)
    sans_agence = 0
    for agence_id, contrats in par_agence:
        service_id = services.get(agence_id) if agence_id is not None else None
        if service_id is None:
            sans_agence += contrats
            continue
        par_service[service_id] += contrats
    await repo.remplacer_nombre_client_mois(db, annee, mois, dict(par_service))
    clients_postes = await repo.clients_par_poste(db, annee, mois)
    await repo.mettre_a_jour_clients_postes(db, clients_postes)
    return BilanImport(
        annee=annee,
        mois=mois,
        source="",
        unites_releve=0,
        contrats=sum(n for _, n in par_agence),
        contrats_sans_agence=sans_agence,
        postes_mis_a_jour=len(clients_postes),
        clients_postes=sum(clients_postes.values()),
        services=len(par_service),
        lignes_rejetees=0,
    )


async def synchroniser_bi_ods(db: AsyncSession, maintenant: datetime | None = None) -> BilanImport:
    """Photo du mois de mise à jour BI_ODS (le dernier chargement du mois fait foi)."""
    brutes = await lire_bi_ods()
    dates = [b["date_maj"] for b in brutes if isinstance(b.get("date_maj"), datetime)]
    ref = max(dates) if dates else (maintenant or datetime.now())
    return await importer(db, ref.year, ref.month, brutes, SOURCE_BI)


# ── API ──────────────────────────────────────────────────────────────────────

_EDITEURS = frozenset({"ADMIN_SCR", "ADMIN_SAQS"})
_TAILLE_MAX = 30 * 1024 * 1024


def peut_administrer(user: Any) -> bool:
    from app.core.rbac_helpers import is_superadmin

    return is_superadmin(user) or any(ur.role.nom in _EDITEURS for ur in user.user_roles)


async def resume(db: AsyncSession) -> dict[str, Any]:
    return {
        "periodes": await repo.resume_periodes(db),
        "liens": await repo.compter_liens(db),
        "bi_ods_configure": bool(settings.bi_ods_odbc),
    }


def lancer_synchronisation() -> str:
    from app.workers.celery_app import celery_app

    if not settings.bi_ods_odbc:
        raise BaseClientsErreur("Connexion BI_ODS non configurée (BI_ODS_ODBC) : déposez le fichier mensuel")
    return celery_app.send_task("synchroniser_contrats_ur").id


def lancer_import_fichier(contenu: bytes, annee: int, mois: int) -> str:
    """Enregistre le fichier déposé et confie l'import au worker (52 000 lignes environ)."""
    import uuid

    from app.workers.celery_app import celery_app

    if not (1 <= mois <= 12 and 2000 <= annee <= 2100):
        raise BaseClientsErreur("Période invalide")
    if len(contenu) > _TAILLE_MAX:
        raise BaseClientsErreur("Fichier trop volumineux (30 Mo au plus)")
    if not contenu.startswith(b"PK"):
        raise BaseClientsErreur("Fichier Excel .xlsx attendu")
    dossier = Path(settings.rapports_storage_path).parent / "imports"
    dossier.mkdir(parents=True, exist_ok=True)
    chemin = dossier / f"contrats_{annee}{mois:02d}_{uuid.uuid4().hex}.xlsx"
    chemin.write_bytes(contenu)
    return celery_app.send_task("importer_contrats_fichier", args=[str(chemin), annee, mois]).id
