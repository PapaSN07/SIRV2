"""Fichiers du rapport hebdomadaire DD : Word (modèle « Rapport Hebdo DD »), Excel (tableau de
bord jour par jour, modèle « TB suivi qualité réseau HTA ») et PDF (même contenu que le Word)."""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from app.services.rapport_hebdo_dd_service import (
    CATEGORIES_HTA,
    CATEGORIES_HTB,
    DAKAR,
    JOURS,
    REGIONS,
    Compte,
    RapportHebdoData,
    ecart_pct,
    fmt_date_longue,
    fmt_duree_h,
    fmt_mwh,
    fmt_nombre,
    titre_semaine,
)

_LOGO = Path(__file__).resolve().parent.parent / "assets" / "senelec_logo.png"
_ENTETE = ("Direction de la Distribution", "Département Conduite et Gestion Technique", "Service Conduite des Réseaux")
_TITRE = "Rapport Hebdomadaire d'Exploitation du Réseau de Distribution"


# --------------------------------------------------------------------------------------
# Tableaux communs (Word et PDF)
# --------------------------------------------------------------------------------------

def tableau_repartition(data: RapportHebdoData) -> list[list[str]]:
    s = data.semaine
    services = data.services
    lignes = [["", *services, "TOTAL"]]
    for cle, libelle in CATEGORIES_HTA:
        valeurs = [s.par_service.get((cle, code), Compte()).nombre for code in services]
        lignes.append([libelle, *[str(v) for v in valeurs], str(sum(valeurs))])
    end = [s.end_service.get(code, 0.0) * 1000 for code in services]
    lignes.append(["**END (kWh)**", *[f"**{fmt_nombre(v)}**" for v in end], f"**{fmt_nombre(sum(end))}**"])
    return lignes


def _observation(libelle: str, ecart: float | None) -> str:
    if ecart is None:
        return "—"
    if abs(ecart) < 0.05:
        return "Stable"
    return f"{'Hausse' if ecart > 0 else 'Baisse'} {libelle}"


def _fleche(ecart: float | None) -> str:
    if ecart is None:
        return "—"
    return f"{ecart:+.2f} %".replace(".", ",") + (" ↗" if ecart > 0 else " ↘" if ecart < 0 else "")


def tableau_comparaison(data: RapportHebdoData) -> list[list[str]]:
    s, p = data.semaine, data.precedente
    lignes = [["", f"Réalisation S{s.numero:02d}", f"Réalisation S{p.numero:02d}", "Écart (%)", "Observations"]]
    elements = (
        ("Interruptions HTA", s.total_hta.nombre, p.total_hta.nombre, str, "du nombre d'interruptions"),
        ("END (MWh)", s.total_hta.end_mwh, p.total_hta.end_mwh, fmt_mwh, "de l'END"),
        ("SAIFI", s.saifi, p.saifi, lambda v: fmt_nombre(v, 3), "de la fréquence moyenne"),
        ("SAIDI (h:mm:ss)", s.saidi_h, p.saidi_h, fmt_duree_h, "de la durée moyenne"),
    )
    for libelle, va, vp, fmt, obs in elements:
        ecart = ecart_pct(float(va), float(vp))
        lignes.append([libelle, fmt(va), fmt(vp), _fleche(ecart), _observation(obs, ecart)])
    return lignes


def tableau_saifi_saidi(saifi: float, saidi_h: float, impactes: int, reference: int) -> list[list[str]]:
    return [
        ["SAIFI", "SAIDI", "Nombre de clients impactés", "Nombre total de clients"],
        [fmt_nombre(saifi, 3), fmt_duree_h(saidi_h), fmt_nombre(impactes), fmt_nombre(reference)],
    ]


def _cumul_lignes(data: RapportHebdoData) -> list[list[str]]:
    fin = fmt_date_longue(data.semaine.fin)
    cible = data.cible_cumul_mwh
    return [
        [f"Cumul END au {fin}", "Valeur cible au " + fin],
        [f"{fmt_nombre(data.cumul_end_mwh / 1000, 3)} GWh", f"{fmt_nombre(cible / 1000, 3)} GWh" if cible else "Objectif non saisi"],
    ]


def _exogenes_lignes(data: RapportHebdoData) -> list[list[str]]:
    return [
        ["Cumul END due aux facteurs exogènes", "Cumul nombre d'interruptions (agressions, chocs…)"],
        [f"{fmt_mwh(data.cumul_end_exogenes_mwh)} MWh", str(data.cumul_nombre_exogenes)],
    ]


def _note_bas() -> str:
    return (
        f"Rapport généré automatiquement par SIR le {datetime.now():%d/%m/%Y à %Hh%M}. "
        "Données HTA de SIR (toutes interruptions apparues dans la semaine, quel que soit leur statut). "
        "Partie BT non incluse."
    )


# --------------------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------------------

def exporter_docx(chemin: Path, data: RapportHebdoData) -> None:
    from app.utils.docx_simple import Docx

    doc = Docx()
    if _LOGO.is_file():
        doc.image(_LOGO.read_bytes(), 3.2, 2.26)
    for i, ligne in enumerate(_ENTETE):
        doc.paragraphe(ligne, gras=i == 0, apres=0)
    doc.paragraphe("")
    doc.paragraphe(_TITRE, gras=True, taille=15, couleur="5B2C6F", align="center", apres=60)
    doc.paragraphe(f"Semaine {titre_semaine(data)}", align="center", taille=11, apres=240)

    doc.paragraphe(
        f"Trouvez ci-joint le rapport hebdomadaire d'exploitation de la Direction de la Distribution "
        f"de la semaine S{data.semaine.numero:02d}-{data.semaine.fin.year}."
    )
    doc.paragraphe(data.introduction)
    doc.paragraphe("**Les faits marquants de la semaine :**")
    doc.paragraphe("**Faits majeurs :**", apres=60)
    doc.puces(data.faits.majeurs or ["RAS"])
    doc.paragraphe("**DAKAR :**" + ("" if data.faits.dakar else " RAS"), apres=60)
    doc.puces(data.faits.dakar)
    doc.paragraphe("**RÉGIONS :**" + ("" if data.faits.regions else " RAS"), apres=60)
    doc.puces(data.faits.regions)
    for paragraphe in data.paragraphes_semaines:
        doc.paragraphe(paragraphe)

    doc.titre("1. Répartition des interruptions par service de distribution", 2)
    nb = len(data.services) + 1
    largeur = max(1.1, 12.3 / nb)
    doc.tableau(tableau_repartition(data), largeurs_cm=[4.7, *[largeur] * nb], derniere_ligne_grasse=True, taille=8)
    if data.dominance:
        doc.paragraphe(data.dominance)
    ends = [(code, data.semaine.end_service.get(code, 0.0)) for code in data.services]
    ends = [(c, v) for c, v in ends if v > 0]
    if ends:
        doc.camembert("END par service (MWh)", [c for c, _ in ends], [v for _, v in ends])

    doc.titre(f"1.1. Cumul de l'END en {data.semaine.fin.year} (au {fmt_date_longue(data.semaine.fin)})", 3)
    doc.tableau(_cumul_lignes(data), largeurs_cm=[8.5, 8.5], premiere_gauche=False)
    doc.titre(f"1.2. Cumul END facteurs exogènes (au {fmt_date_longue(data.semaine.fin)})", 3)
    doc.tableau(_exogenes_lignes(data), largeurs_cm=[8.5, 8.5], premiere_gauche=False)
    doc.titre(
        f"1.3. Comparaison des réalisations entre S{data.semaine.numero:02d} et S{data.precedente.numero:02d}", 3
    )
    doc.tableau(tableau_comparaison(data), largeurs_cm=[3.6, 2.9, 2.9, 2.8, 4.8])
    doc.titre("1.4. Éléments SAIFI et SAIDI", 3)
    doc.paragraphe(f"**Semaine du {fmt_date_longue(data.semaine.debut)} au {fmt_date_longue(data.semaine.fin)}**", apres=60)
    s = data.semaine
    doc.tableau(
        tableau_saifi_saidi(s.saifi, s.saidi_h, s.clients_impactes, s.clients_reference),
        largeurs_cm=[3.5, 3.5, 5, 5],
        premiere_gauche=False,
    )
    doc.paragraphe(f"**Cumul {s.fin.year} (jusqu'au {fmt_date_longue(s.fin)})**", apres=60)
    doc.tableau(
        tableau_saifi_saidi(data.cumul_saifi, data.cumul_saidi_h, data.cumul_clients_impactes, data.cumul_clients_reference),
        largeurs_cm=[3.5, 3.5, 5, 5],
        premiere_gauche=False,
    )
    doc.paragraphe(_note_bas(), taille=8, couleur="7F7F7F")
    doc.enregistrer(chemin)


# --------------------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------------------

def exporter_pdf(chemin: Path, data: RapportHebdoData) -> None:
    import re

    from reportlab.graphics.charts.legends import Legend
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import Image, KeepTogether, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    violet = colors.HexColor("#5B2C6F")
    styles = getSampleStyleSheet()
    normal = ParagraphStyle("n", parent=styles["Normal"], fontSize=9.5, leading=13, spaceAfter=5)
    petit = ParagraphStyle("p", parent=normal, fontSize=7.5, leading=9, textColor=colors.HexColor("#7F7F7F"))
    titre = ParagraphStyle("t", parent=normal, fontSize=15, leading=19, alignment=TA_CENTER, textColor=violet,
                           fontName="Helvetica-Bold")
    sous_titre = ParagraphStyle("st", parent=normal, alignment=TA_CENTER, fontSize=10.5)
    h2 = ParagraphStyle("h2", parent=normal, fontSize=12, leading=15, textColor=violet, fontName="Helvetica-Bold",
                        spaceBefore=8)
    h3 = ParagraphStyle("h3", parent=h2, fontSize=10.5, leading=13)
    puce = ParagraphStyle("pu", parent=normal, leftIndent=14, bulletIndent=4, spaceAfter=2)
    cellule = ParagraphStyle("c", parent=normal, fontSize=7.5, leading=9, spaceAfter=0, alignment=TA_CENTER)
    cellule_g = ParagraphStyle("cg", parent=cellule, alignment=0)

    def p(texte: str, style: ParagraphStyle = normal) -> Paragraph:
        texte = texte.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(re.sub(r"\*\*([^*]+)\*\*", r"<b>\1</b>", texte), style)

    def table(
        lignes: list[list[str]], largeurs: list[float], gras_fin: bool = False, premiere_gauche: bool = True
    ) -> Table:
        contenu = []
        for i, ligne in enumerate(lignes):
            cellules = []
            for j, v in enumerate(ligne):
                v = v.replace(" ↗", "").replace(" ↘", "")  # flèches absentes de la police Helvetica
                if i == 0 and v:
                    v = f"**{v}**"
                cellules.append(p(v, cellule_g if (j == 0 and premiere_gauche) else cellule))
            contenu.append(cellules)
        t = Table(contenu, colWidths=[w * cm for w in largeurs], repeatRows=1)
        style = [
            ("GRID", (0, 0), (-1, -1), 0.4, colors.HexColor("#BFBFBF")),
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EDE3F2")),
            ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
            ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
            ("LEFTPADDING", (0, 0), (-1, -1), 3),
            ("RIGHTPADDING", (0, 0), (-1, -1), 3),
        ]
        if gras_fin:
            style.append(("BACKGROUND", (0, -1), (-1, -1), colors.HexColor("#F6F1F8")))
        t.setStyle(TableStyle(style))
        return t

    elements: list = []
    if _LOGO.is_file():
        elements.append(Image(str(_LOGO), width=3.2 * cm, height=2.26 * cm, hAlign="LEFT"))
    for i, ligne in enumerate(_ENTETE):
        elements.append(p(f"**{ligne}**" if i == 0 else ligne, ParagraphStyle("e", parent=normal, spaceAfter=0)))
    elements += [Spacer(1, 10), p(_TITRE, titre), p(f"Semaine {titre_semaine(data)}", sous_titre), Spacer(1, 8)]
    elements.append(
        p(
            f"Trouvez ci-joint le rapport hebdomadaire d'exploitation de la Direction de la Distribution de la "
            f"semaine S{data.semaine.numero:02d}-{data.semaine.fin.year}."
        )
    )
    elements.append(p(data.introduction))
    elements.append(p("**Les faits marquants de la semaine :**"))
    elements.append(p("**Faits majeurs :**"))
    elements += [p(f"•  {f}", puce) for f in (data.faits.majeurs or ["RAS"])]
    elements.append(p("**DAKAR :**" + ("" if data.faits.dakar else " RAS")))
    elements += [p(f"•  {f}", puce) for f in data.faits.dakar]
    elements.append(p("**RÉGIONS :**" + ("" if data.faits.regions else " RAS")))
    elements += [p(f"•  {f}", puce) for f in data.faits.regions]
    for paragraphe in data.paragraphes_semaines:
        elements.append(p(paragraphe))

    elements.append(p("1. Répartition des interruptions par service de distribution", h2))
    nb = len(data.services) + 1
    largeur = max(1.0, 13.0 / nb)
    elements.append(table(tableau_repartition(data), [4.6, *[largeur] * nb], gras_fin=True))
    if data.dominance:
        elements += [Spacer(1, 4), p(data.dominance)]
    ends = [(c, data.semaine.end_service.get(c, 0.0)) for c in data.services]
    ends = [(c, v) for c, v in ends if v > 0]
    if ends:
        dessin = Drawing(16 * cm, 6.5 * cm)
        pie = Pie()
        pie.x, pie.y, pie.width, pie.height = 1.5 * cm, 0.5 * cm, 5.5 * cm, 5.5 * cm
        pie.data = [v for _, v in ends]
        total = sum(pie.data)
        pie.labels = [f"{v * 100 / total:.0f} %" for v in pie.data]
        palette = ["#5B2C6F", "#E4007C", "#F39200", "#C0392B", "#2E86C1", "#16A085", "#8E44AD", "#D35400",
                   "#7F8C8D", "#27AE60", "#B03A2E", "#1F618D"]
        for i in range(len(pie.data)):
            pie.slices[i].fillColor = colors.HexColor(palette[i % len(palette)])
            pie.slices[i].strokeColor = colors.white
        pie.slices.fontSize = 7
        pie.slices.fontName = "Helvetica"
        dessin.add(pie)
        legende = Legend()
        legende.x, legende.y = 8.5 * cm, 5.8 * cm
        legende.alignment = "right"
        legende.columnMaximum = len(ends)
        legende.fontName = "Helvetica"
        legende.fontSize = 8
        legende.dy = legende.dx = 6
        legende.deltay = 11
        legende.colorNamePairs = [
            (colors.HexColor(palette[i % len(palette)]), f"{c} — {fmt_mwh(v)} MWh") for i, (c, v) in enumerate(ends)
        ]
        dessin.add(legende)
        elements.append(dessin)

    fin = data.semaine.fin
    elements.append(KeepTogether([p(f"1.1. Cumul de l'END en {fin.year} (au {fmt_date_longue(fin)})", h3),
                                  table(_cumul_lignes(data), [8.5, 8.5], premiere_gauche=False)]))
    elements.append(KeepTogether([p(f"1.2. Cumul END facteurs exogènes (au {fmt_date_longue(fin)})", h3),
                                  table(_exogenes_lignes(data), [8.5, 8.5], premiere_gauche=False)]))
    elements.append(KeepTogether([
        p(f"1.3. Comparaison des réalisations entre S{data.semaine.numero:02d} et S{data.precedente.numero:02d}", h3),
        table(tableau_comparaison(data), [3.6, 2.9, 2.9, 2.8, 4.8]),
    ]))
    s = data.semaine
    elements.append(KeepTogether([
        p("1.4. Éléments SAIFI et SAIDI", h3),
        p(f"**Semaine du {fmt_date_longue(s.debut)} au {fmt_date_longue(s.fin)}**"),
        table(tableau_saifi_saidi(s.saifi, s.saidi_h, s.clients_impactes, s.clients_reference), [3.5, 3.5, 5, 5],
              premiere_gauche=False),
        Spacer(1, 6),
        p(f"**Cumul {fin.year} (jusqu'au {fmt_date_longue(fin)})**"),
        table(
            tableau_saifi_saidi(data.cumul_saifi, data.cumul_saidi_h, data.cumul_clients_impactes, data.cumul_clients_reference),
            [3.5, 3.5, 5, 5],
            premiere_gauche=False,
        ),
    ]))
    elements += [Spacer(1, 10), p(_note_bas(), petit)]

    doc = SimpleDocTemplate(
        str(chemin), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm, topMargin=1.6 * cm, bottomMargin=1.6 * cm,
        title=f"Rapport hebdomadaire DD {titre_semaine(data)}", author="SIR — SENELEC",
    )
    doc.build(elements)


# --------------------------------------------------------------------------------------
# Excel — tableau de bord jour par jour
# --------------------------------------------------------------------------------------

def exporter_xlsx(chemin: Path, data: RapportHebdoData) -> None:
    from openpyxl import Workbook
    from openpyxl.chart import PieChart, Reference
    from openpyxl.chart.label import DataLabelList
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    violet = "5B2C6F"
    fond_entete = PatternFill("solid", fgColor="EDE3F2")
    fond_total = PatternFill("solid", fgColor="F6F1F8")
    fin_bord = Side(style="thin", color="BFBFBF")
    bord = Border(left=fin_bord, right=fin_bord, top=fin_bord, bottom=fin_bord)
    gras = Font(bold=True)
    centre = Alignment(horizontal="center", vertical="center", wrap_text=True)
    s = data.semaine
    jours = [s.debut.toordinal() + i for i in range(7)]

    wb = Workbook()
    ws = wb.active
    ws.title = "Synthèse quotidienne"

    def entete_feuille(feuille, titre: str) -> None:
        feuille["A1"] = "Direction de la Distribution"
        feuille["A1"].font = Font(bold=True, color=violet, size=12)
        feuille["A2"] = titre
        feuille["A2"].font = Font(bold=True, size=11)
        feuille["A3"] = f"Semaine {titre_semaine(data)}"

    def cellule(
        feuille, ligne: int, col: int, valeur, *, en_gras: bool = False, fond=None, fmt: str | None = None,
        gauche: bool = False,
    ):
        c = feuille.cell(row=ligne, column=col, value=valeur)
        c.border = bord
        c.alignment = Alignment(vertical="center", wrap_text=True) if (col == 1 or gauche) else centre
        if en_gras:
            c.font = gras
        if fond:
            c.fill = fond
        if fmt:
            c.number_format = fmt
        return c

    entete_feuille(ws, "Tableau de bord de suivi qualité des réseaux de distribution — paramètres du réseau HTA")
    ligne = 5
    cellule(ws, ligne, 1, "Zone", en_gras=True, fond=fond_entete)
    cellule(ws, ligne, 2, "Paramètres", en_gras=True, fond=fond_entete)
    ws.merge_cells(start_row=ligne, start_column=3, end_row=ligne, end_column=10)
    cellule(ws, ligne, 3, "Nombre", en_gras=True, fond=fond_entete)
    ws.merge_cells(start_row=ligne, start_column=11, end_row=ligne, end_column=18)
    cellule(ws, ligne, 11, "END (kWh)", en_gras=True, fond=fond_entete)
    ligne += 1
    cellule(ws, ligne, 1, "", fond=fond_entete)
    cellule(ws, ligne, 2, "", fond=fond_entete)
    for bloc in (3, 11):
        for i, j in enumerate(jours):
            jour = date.fromordinal(j)
            cellule(ws, ligne, bloc + i, f"{JOURS[i][:3]}. {jour:%d/%m}", en_gras=True, fond=fond_entete)
        cellule(ws, ligne, bloc + 7, "Cumul semaine", en_gras=True, fond=fond_entete)
    ligne += 1

    categories = [*CATEGORIES_HTA, ("exogenes", "dont facteurs exogènes (choc accidentel…)"), *CATEGORIES_HTB]
    for z in (DAKAR, REGIONS, None):
        debut_bloc = ligne
        for cle, libelle in categories:
            cellule(ws, ligne, 1, z or "TOTAL", en_gras=True)
            cellule(ws, ligne, 2, libelle, en_gras=cle == "exogenes", gauche=True)
            for i in range(7):
                if cle == "exogenes":
                    comptes = [s.exogenes_jour.get((i, zz), Compte()) for zz in ((z,) if z else (DAKAR, REGIONS))]
                else:
                    comptes = [s.par_jour.get((i, cle, zz), Compte()) for zz in ((z,) if z else (DAKAR, REGIONS))]
                cellule(ws, ligne, 3 + i, sum(c.nombre for c in comptes))
                cellule(ws, ligne, 11 + i, round(sum(c.end_mwh for c in comptes) * 1000), fmt="# ##0")
            cellule(ws, ligne, 10, f"=SUM({get_column_letter(3)}{ligne}:{get_column_letter(9)}{ligne})", en_gras=True)
            cellule(ws, ligne, 18, f"=SUM({get_column_letter(11)}{ligne}:{get_column_letter(17)}{ligne})",
                    en_gras=True, fmt="# ##0")
            ligne += 1
        ws.merge_cells(start_row=debut_bloc, start_column=1, end_row=ligne - 1, end_column=1)
        nb_hta = len(CATEGORIES_HTA)
        cellule(ws, ligne, 1, f"Total HTA {z or 'national'}", en_gras=True, fond=fond_total)
        cellule(ws, ligne, 2, "Toutes interruptions HTA", en_gras=True, fond=fond_total, gauche=True)
        for col in list(range(3, 11)) + list(range(11, 19)):
            lettre = get_column_letter(col)
            cellule(ws, ligne, col, f"=SUM({lettre}{debut_bloc}:{lettre}{debut_bloc + nb_hta - 1})", en_gras=True,
                    fond=fond_total, fmt="# ##0" if col >= 11 else None)
        ligne += 2

    ws.column_dimensions["A"].width = 14
    ws.column_dimensions["B"].width = 44
    for col in range(3, 19):
        ws.column_dimensions[get_column_letter(col)].width = 10.5
    ws.freeze_panes = "C7"

    # ---- Répartition par service --------------------------------------------------------
    wr = wb.create_sheet("Répartition par SD")
    entete_feuille(wr, "Répartition des interruptions par service de distribution")
    services = data.services
    ligne = 5
    blocs: list[tuple[str, int | None]] = [(f"{JOURS[i]} {date.fromordinal(jours[i]):%d/%m/%Y}", i) for i in range(7)]
    blocs.append((f"Semaine S{s.numero:02d}", None))
    ligne_end_semaine = 0
    for titre_bloc, jour in blocs:
        cellule(wr, ligne, 1, titre_bloc, en_gras=True, fond=fond_entete)
        for j, code in enumerate(services):
            cellule(wr, ligne, 2 + j, code, en_gras=True, fond=fond_entete)
        cellule(wr, ligne, 2 + len(services), "TOTAL", en_gras=True, fond=fond_entete)
        ligne += 1
        for cle, libelle in CATEGORIES_HTA:
            cellule(wr, ligne, 1, libelle)
            for j, code in enumerate(services):
                c = s.par_service.get((cle, code), Compte()) if jour is None else s.par_jour_service.get((jour, cle, code), Compte())
                cellule(wr, ligne, 2 + j, c.nombre)
            fin_col = get_column_letter(1 + len(services))
            cellule(wr, ligne, 2 + len(services), f"=SUM(B{ligne}:{fin_col}{ligne})", en_gras=True)
            ligne += 1
        cellule(wr, ligne, 1, "END (kWh)", en_gras=True, fond=fond_total)
        for j, code in enumerate(services):
            if jour is None:
                valeur = s.end_service.get(code, 0.0)
            else:
                valeur = sum(s.par_jour_service.get((jour, cle, code), Compte()).end_mwh for cle, _ in CATEGORIES_HTA)
            cellule(wr, ligne, 2 + j, round(valeur * 1000), en_gras=True, fond=fond_total, fmt="# ##0")
        fin_col = get_column_letter(1 + len(services))
        cellule(wr, ligne, 2 + len(services), f"=SUM(B{ligne}:{fin_col}{ligne})", en_gras=True, fond=fond_total,
                fmt="# ##0")
        if jour is None:
            ligne_end_semaine = ligne
        ligne += 2
    wr.column_dimensions["A"].width = 40
    for j in range(len(services) + 1):
        wr.column_dimensions[get_column_letter(2 + j)].width = 10
    if services and ligne_end_semaine:
        camembert = PieChart()
        camembert.title = f"END par service — S{s.numero:02d}"
        valeurs = Reference(wr, min_col=2, max_col=1 + len(services), min_row=ligne_end_semaine)
        libelles = Reference(wr, min_col=2, max_col=1 + len(services), min_row=ligne_end_semaine - len(CATEGORIES_HTA) - 1)
        camembert.add_data(valeurs, from_rows=True, titles_from_data=False)
        camembert.set_categories(libelles)
        camembert.dataLabels = DataLabelList()
        camembert.dataLabels.showPercent = True
        camembert.dataLabels.showVal = False
        camembert.dataLabels.showCatName = False
        camembert.dataLabels.showSerName = False
        camembert.dataLabels.showLegendKey = False
        camembert.height, camembert.width = 9, 14
        wr.add_chart(camembert, f"{get_column_letter(len(services) + 4)}5")

    # ---- Synthèse (mêmes chiffres que le Word) -----------------------------------------
    wsy = wb.create_sheet("Synthèse semaine")
    entete_feuille(wsy, "Synthèse de la semaine")
    lignes_synthese = [
        ("Interruptions HTA (semaine)", s.total_hta.nombre),
        ("END HTA semaine (MWh)", round(s.total_hta.end_mwh, 3)),
        ("dont facteurs exogènes (MWh)", round(s.exogenes.end_mwh, 3)),
        ("SAIFI semaine", round(s.saifi, 4)),
        ("SAIDI semaine (h)", round(s.saidi_h, 4)),
        ("Clients impactés (semaine)", s.clients_impactes),
        ("Nombre total de clients", s.clients_reference),
        (f"Cumul END {s.fin.year} (MWh)", round(data.cumul_end_mwh, 3)),
        ("Valeur cible cumulée (MWh)", round(data.cible_cumul_mwh, 3) if data.cible_cumul_mwh else "Objectif non saisi"),
        (f"Cumul END exogènes {s.fin.year} (MWh)", round(data.cumul_end_exogenes_mwh, 3)),
        ("Cumul interruptions exogènes", data.cumul_nombre_exogenes),
        (f"SAIFI cumul {s.fin.year}", round(data.cumul_saifi, 4)),
        (f"SAIDI cumul {s.fin.year} (h)", round(data.cumul_saidi_h, 4)),
    ]
    for i, (libelle, valeur) in enumerate(lignes_synthese, start=5):
        cellule(wsy, i, 1, libelle, en_gras=True)
        cellule(wsy, i, 2, valeur)
    base = 5 + len(lignes_synthese) + 1
    wsy.cell(row=base, column=1, value="Faits marquants").font = Font(bold=True, color=violet)
    r = base + 1
    for titre_section, faits in (("Faits majeurs", data.faits.majeurs), ("Dakar", data.faits.dakar), ("Régions", data.faits.regions)):
        wsy.cell(row=r, column=1, value=titre_section).font = gras
        r += 1
        for f in faits or ["RAS"]:
            wsy.cell(row=r, column=1, value=f"• {f}")
            r += 1
    wsy.column_dimensions["A"].width = 46
    wsy.column_dimensions["B"].width = 18

    for feuille in (ws, wr, wsy):
        feuille.page_setup.orientation = "landscape"
        feuille.page_setup.paperSize = feuille.PAPERSIZE_A4
        feuille.page_setup.fitToWidth = 1
        feuille.page_setup.fitToHeight = 0
        feuille.sheet_properties.pageSetUpPr.fitToPage = True
    wb.save(chemin)


def exporter_tout(dossier: Path, nom_base: str, data: RapportHebdoData) -> dict[str, Path]:
    """Écrit les trois fichiers ; renvoie {extension: chemin}."""
    chemins = {ext: dossier / f"{nom_base}.{ext}" for ext in ("docx", "xlsx", "pdf")}
    exporter_docx(chemins["docx"], data)
    exporter_xlsx(chemins["xlsx"], data)
    exporter_pdf(chemins["pdf"], data)
    return chemins
