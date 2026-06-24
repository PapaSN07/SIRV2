"""Service import FIS (Fiche d'Incidents et de Suivi) — BCC et BRC.

Lit le fichier Excel FIS (format SENELEC), extrait la feuille
«Détails interruptions» et crée les interruptions distribution en statut
«valide» avec l'END déjà calculé (valeur issue du fichier FIS).

Formule END vérifiée sur données réelles :
    END (kWh) = Puissance_coupée (kW) × Durée (minutes) / 60

Correspondance FIS → SIRV2 :
    END_mwh = END_kwh / 1000
    puissance_mw = puissance_kw / 1000
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from typing import Any

from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.models.geo import Service
from app.models.interruptions import Interruption, InterruptionDistribution
from app.models.referentiel import Cause, Nature, Siege
from app.models.reseau import Feeder, PosteSource
from app.models.utilisateurs import Utilisateur

logger = logging.getLogger(__name__)

# ---------------------------------------------------------------------------
# Constantes de mapping FIS → codes référentiel SIRV2
# ---------------------------------------------------------------------------

# Nature présumée FIS libelle → code SIRV2
NATURE_LIBELLE_TO_CODE: dict[str, str] = {
    "permanent": "PERMANENT",
    "semi permanent": "SEMI_PERM",
    "fugitif": "FUGITIF",
    "délestage": "DELESTAGE",
    "coupure": "COUPURE",
}

# Cause présumée FIS libelle → (code SIRV2, is_force_majeure)
CAUSE_LIBELLE_TO_CODE: dict[str, tuple[str, bool]] = {
    "inconnue": ("INCONNUE", False),
    "agression": ("AGRESSION", True),
    "incident réseau hta": ("INC_HTA", False),
    "défaillance protections dt": ("DEF_PROT", False),
    "défaillance matériel dt": ("DEF_MAT_DT", False),
    "défaillance matériel htb": ("DEF_MAT_HTB", False),
    "travaux réseau hta": ("TRAV_HTA", False),
    "travaux réseau htb": ("TRAV_HTB", False),
    "travaux réseau bt": ("TRAV_BT", False),
    "facteurs exogènes": ("FACT_EXO", True),
    "incendie": ("INCENDIE", True),
    "non recherché": ("NON_RECHERCHE", False),
    "manoeuvre": ("MANOEUVRE", False),
    "essai sur défaut": ("ESSAI_DEF", False),
    "défaillance protection htb": ("DEF_PROT_HTB", False),
    "manque production": ("MANQUE_PROD", False),
}

# Siège présumé FIS libelle → code SIRV2
SIEGE_LIBELLE_TO_CODE: dict[str, str] = {
    "câble": "CABLE",
    "extrémité": "EXTREMITE",
    "ligature": "LIGATURE",
    "ligne": "LIGNE",
    "cellule": "CELLULE",
    "disjoncteur": "DISJONCTEUR",
    "transformateur": "TRANSFO",
    "autre élément hta": "AUTRE_HTA",
    "autre": "AUTRE",
}

# Natures qui correspondent à des coupures programmées (is_programme=True)
NATURES_PROGRAMMEES = {"COUPURE"}

# Natures fugitives
NATURES_FUGITIVES = {"FUGITIF"}

# Natures délestage (is_programme=False mais délestage)
NATURES_DELESTAGE = {"DELESTAGE"}

# Feuille contenant les données d'incidents
SHEET_NAME = "Détails interruptions"

# En-têtes attendus dans la feuille (insensible à la casse)
_COL_NUMERO = "n°"
_COL_SERVICE = "service distribution"
_COL_SOURCE = "source"
_COL_TENSION = "tension primaire"
_COL_FEEDER = "feeder"
_COL_NATURE = "nature présumée"
_COL_CAUSE = "cause présumée"
_COL_SIEGE = "siege présumé"
_COL_SIGNALISATION = "signalisation"
_COL_DATE_APPARITION = "date d'apparition"
_COL_DATE_REMISE = "date de remise"
_COL_DUREE = "durée interruption"
_COL_PUISSANCE = "puissance coupée"
_COL_END = "end"


# ---------------------------------------------------------------------------
# Structures de résultats
# ---------------------------------------------------------------------------

@dataclass
class FisRow:
    """Ligne parsée depuis la feuille 'Détails interruptions'."""
    numero: int
    service_libelle: str
    source_libelle: str
    tension_kv: float | None
    feeder_libelle: str
    nature_libelle: str
    cause_libelle: str
    siege_libelle: str | None
    date_apparition: datetime
    date_remise: datetime | None
    duree_minutes: int
    puissance_kw: float
    end_kwh: float
    troncon: str | None = None
    observations: str | None = None


@dataclass
class FisImportStats:
    total_lignes: int = 0
    importees: int = 0
    doublons_ignores: int = 0
    erreurs: int = 0
    end_total_kwh: float = 0.0
    erreurs_detail: list[str] = field(default_factory=list)


@dataclass
class FisImportResult:
    stats: FisImportStats
    periode_debut: datetime | None = None
    periode_fin: datetime | None = None
    services_touches: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Parsing Excel
# ---------------------------------------------------------------------------

def _normalise(s: str | None) -> str:
    """Normalise une chaîne : strip + lowercase."""
    if s is None:
        return ""
    return s.strip().lower()


def _find_header_row(ws) -> dict[str, int] | None:  # type: ignore[type-arg]
    """Détecte la ligne d'en-tête et retourne un dict col_name → col_index."""
    for row in ws.iter_rows(values_only=True):
        # La ligne d'en-tête contient "N°" ou "n°" dans la 1re cellule non-None
        cells = [c for c in row if c is not None]
        if not cells:
            continue
        first = _normalise(str(cells[0]))
        if first == _COL_NUMERO:
            return {
                _normalise(str(v)): i
                for i, v in enumerate(row)
                if v is not None
            }
    return None


def _col(headers: dict[str, int], name: str) -> int | None:
    """Retourne l'index de colonne pour le nom (recherche préfixe insensible casse)."""
    # Exact match
    if name in headers:
        return headers[name]
    # Prefix match (ex: "source" matche "source ")
    for key, idx in headers.items():
        if key.startswith(name):
            return idx
    return None


def _parse_duration(val: Any) -> int:
    """Convertit une valeur de durée en minutes entiers."""
    if val is None:
        return 0
    if isinstance(val, timedelta):
        return int(val.total_seconds() / 60)
    if isinstance(val, str):
        # Format "0 h 31 m 00" ou "1 h 22 m 00"
        import re
        m = re.match(r"(\d+)\s*h\s*(\d+)\s*m", val)
        if m:
            return int(m.group(1)) * 60 + int(m.group(2))
    return 0


def _parse_float(val: Any) -> float:
    """Convertit une valeur numérique en float."""
    if val is None:
        return 0.0
    if isinstance(val, (int, float)):
        return float(val)
    try:
        return float(str(val).replace(",", "."))
    except (ValueError, TypeError):
        return 0.0


def _parse_datetime(val: Any) -> datetime | None:
    """Retourne un datetime ou None."""
    if isinstance(val, datetime):
        return val
    return None


def parse_fis_sheet(file_bytes: bytes) -> list[FisRow]:
    """Lit un fichier FIS Excel et retourne les lignes de 'Détails interruptions'."""
    try:
        import openpyxl
    except ImportError as exc:
        raise RuntimeError("openpyxl requis : pip install openpyxl") from exc

    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), read_only=True, data_only=True)

    if SHEET_NAME not in wb.sheetnames:
        raise ValueError(
            f"Feuille '{SHEET_NAME}' introuvable. "
            f"Feuilles disponibles : {wb.sheetnames}"
        )

    ws = wb[SHEET_NAME]
    headers = _find_header_row(ws)
    if headers is None:
        raise ValueError(
            f"Ligne d'en-tête introuvable dans '{SHEET_NAME}' "
            "(colonnes attendues : N°, Service Distribution, …)"
        )

    # Index des colonnes clés
    idx_num = _col(headers, _COL_NUMERO)
    idx_svc = _col(headers, _COL_SERVICE)
    idx_src = _col(headers, _COL_SOURCE)
    idx_ten = _col(headers, _COL_TENSION)
    idx_fdr = _col(headers, _COL_FEEDER)
    idx_nat = _col(headers, _COL_NATURE)
    idx_cau = _col(headers, _COL_CAUSE)
    idx_sie = _col(headers, _COL_SIEGE)
    idx_da = _col(headers, _COL_DATE_APPARITION)
    idx_dr = _col(headers, _COL_DATE_REMISE)
    idx_dur = _col(headers, _COL_DUREE)
    idx_pui = _col(headers, _COL_PUISSANCE)
    idx_end = _col(headers, _COL_END)

    # Colonnes optionnelles
    idx_trc = _col(headers, "tronçon")
    idx_obs = _col(headers, "observation")

    required_cols = {
        "N°": idx_num,
        "Service Distribution": idx_svc,
        "Feeder": idx_fdr,
        "Date d'apparition": idx_da,
        "Puissance coupée": idx_pui,
        "END": idx_end,
    }
    missing = [name for name, idx in required_cols.items() if idx is None]
    if missing:
        raise ValueError(f"Colonnes manquantes dans FIS : {missing}")

    rows: list[FisRow] = []
    header_passed = False

    for raw_row in ws.iter_rows(values_only=True):
        # Skip jusqu'après la ligne d'en-tête
        if not header_passed:
            cells = [c for c in raw_row if c is not None]
            if cells and _normalise(str(cells[0])) == _COL_NUMERO:
                header_passed = True
            continue

        # Ignorer les lignes vides ou sans numéro
        num_val = raw_row[idx_num] if idx_num is not None else None  # type: ignore[index]
        if num_val is None or not isinstance(num_val, int):
            continue

        svc_val = raw_row[idx_svc] if idx_svc is not None else None  # type: ignore[index]
        da_val = raw_row[idx_da] if idx_da is not None else None  # type: ignore[index]
        dr_val = raw_row[idx_dr] if idx_dr is not None else None  # type: ignore[index]
        dur_val = raw_row[idx_dur] if idx_dur is not None else None  # type: ignore[index]
        pui_val = raw_row[idx_pui] if idx_pui is not None else None  # type: ignore[index]
        end_val = raw_row[idx_end] if idx_end is not None else None  # type: ignore[index]

        da = _parse_datetime(da_val)
        if da is None:
            continue  # date_apparition obligatoire

        rows.append(
            FisRow(
                numero=int(num_val),
                service_libelle=str(svc_val).strip() if svc_val else "",
                source_libelle=str(raw_row[idx_src]).strip() if idx_src is not None and raw_row[idx_src] else "",  # type: ignore[index]
                tension_kv=_parse_float(raw_row[idx_ten]) if idx_ten is not None else None,  # type: ignore[index]
                feeder_libelle=str(raw_row[idx_fdr]).strip() if idx_fdr is not None and raw_row[idx_fdr] else "",  # type: ignore[index]
                nature_libelle=str(raw_row[idx_nat]).strip() if idx_nat is not None and raw_row[idx_nat] else "",  # type: ignore[index]
                cause_libelle=str(raw_row[idx_cau]).strip() if idx_cau is not None and raw_row[idx_cau] else "",  # type: ignore[index]
                siege_libelle=str(raw_row[idx_sie]).strip() if idx_sie is not None and raw_row[idx_sie] else None,  # type: ignore[index]
                date_apparition=da,
                date_remise=_parse_datetime(dr_val),
                duree_minutes=_parse_duration(dur_val),
                puissance_kw=_parse_float(pui_val),
                end_kwh=_parse_float(end_val),
                troncon=str(raw_row[idx_trc]).strip() if idx_trc is not None and raw_row[idx_trc] else None,  # type: ignore[index]
                observations=str(raw_row[idx_obs]).strip() if idx_obs is not None and raw_row[idx_obs] else None,  # type: ignore[index]
            )
        )

    wb.close()
    return rows


# ---------------------------------------------------------------------------
# Résolution des référentiels DB (get-or-create)
# ---------------------------------------------------------------------------

async def _get_or_create_nature(db: AsyncSession, libelle: str) -> Nature | None:
    """Résout ou crée une Nature depuis le libellé FIS."""
    if not libelle:
        return None
    key = _normalise(libelle)
    code = NATURE_LIBELLE_TO_CODE.get(key)
    if code is None:
        # Essai partiel
        for k, c in NATURE_LIBELLE_TO_CODE.items():
            if k in key or key in k:
                code = c
                break
    if code is None:
        code = "AUTRE"

    result = await db.execute(select(Nature).where(Nature.code == code))
    existing = result.scalar_one_or_none()
    if existing:
        return existing

    nature = Nature(code=code, libelle=libelle.strip(), is_active=True)
    db.add(nature)
    await db.flush()
    return nature


async def _get_or_create_cause(
    db: AsyncSession, libelle: str
) -> tuple[Cause | None, bool]:
    """Résout ou crée une Cause depuis le libellé FIS. Retourne (cause, is_force_majeure)."""
    if not libelle:
        return None, False
    key = _normalise(libelle)
    mapping = CAUSE_LIBELLE_TO_CODE.get(key)
    if mapping is None:
        for k, v in CAUSE_LIBELLE_TO_CODE.items():
            if k in key or key in k:
                mapping = v
                break
    if mapping is None:
        mapping = ("INCONNUE", False)

    code, is_fm = mapping
    result = await db.execute(select(Cause).where(Cause.code == code))
    existing = result.scalar_one_or_none()
    if existing:
        return existing, is_fm

    cause = Cause(code=code, libelle=libelle.strip(), type_reseau="distribution", is_active=True)
    db.add(cause)
    await db.flush()
    return cause, is_fm


async def _get_or_create_siege(db: AsyncSession, libelle: str | None) -> Siege | None:
    """Résout ou crée un Siège depuis le libellé FIS."""
    if not libelle:
        return None
    key = _normalise(libelle)
    code = SIEGE_LIBELLE_TO_CODE.get(key)
    if code is None:
        for k, c in SIEGE_LIBELLE_TO_CODE.items():
            if k in key or key in k:
                code = c
                break
    if code is None:
        code = "AUTRE"

    result = await db.execute(select(Siege).where(Siege.code == code))
    existing = result.scalar_one_or_none()
    if existing:
        return existing

    siege = Siege(code=code, libelle=libelle.strip(), type_reseau="distribution", is_active=True)
    db.add(siege)
    await db.flush()
    return siege


async def _find_service(db: AsyncSession, libelle: str) -> Service | None:
    """Cherche un service par code ou libelle."""
    if not libelle:
        return None
    # Par code exact (ex: SDRBD, SDPG, SDRN1…)
    result = await db.execute(select(Service).where(Service.code == libelle.upper()))
    service = result.scalar_one_or_none()
    if service:
        return service
    # Par libelle insensible casse
    result = await db.execute(
        select(Service).where(
            Service.libelle.ilike(f"%{libelle}%")
        )
    )
    return result.scalars().first()


async def _find_feeder(
    db: AsyncSession, libelle: str, service_id: int | None
) -> Feeder | None:
    """Cherche un feeder par libelle dans le service donné ou globalement."""
    if not libelle:
        return None
    query = select(Feeder).where(Feeder.libelle.ilike(f"%{libelle}%"))
    if service_id:
        query = query.where(Feeder.service_id == service_id)
    result = await db.execute(query.limit(1))
    return result.scalar_one_or_none()


async def _find_poste_source(db: AsyncSession, libelle: str) -> PosteSource | None:
    """Cherche un poste source par libelle."""
    if not libelle:
        return None
    result = await db.execute(
        select(PosteSource)
        .where(PosteSource.libelle.ilike(f"%{libelle}%"))
        .limit(1)
    )
    return result.scalar_one_or_none()


async def _find_system_user(db: AsyncSession) -> Utilisateur:
    """Retourne le premier utilisateur actif pour les imports (utilisateur système)."""
    result = await db.execute(
        select(Utilisateur)
        .where(Utilisateur.is_active == True)  # noqa: E712
        .limit(1)
    )
    user = result.scalar_one_or_none()
    if user is None:
        raise RuntimeError(
            "Aucun utilisateur actif trouvé — "
            "exécuter les seeds avant d'importer un fichier FIS"
        )
    return user


async def _interruption_exists(db: AsyncSession, numero: int) -> bool:
    """Vérifie si une interruption avec ce numéro existe déjà."""
    result = await db.execute(
        select(Interruption.id).where(
            Interruption.numero_interruption == numero,
            Interruption.type_direction == "distribution",
        )
    )
    return result.scalar_one_or_none() is not None


# ---------------------------------------------------------------------------
# Import principal
# ---------------------------------------------------------------------------

async def import_fis_file(
    db: AsyncSession,
    file_bytes: bytes,
    *,
    created_by_id: int,
    skip_existing: bool = True,
) -> FisImportResult:
    """Parse le fichier FIS et importe les interruptions en base.

    Args:
        db: session SQLAlchemy async.
        file_bytes: contenu brut du fichier Excel FIS.
        created_by_id: ID de l'utilisateur qui déclenche l'import.
        skip_existing: si True, ignore les numéros déjà présents en base.

    Returns:
        FisImportResult avec statistiques d'import.
    """
    stats = FisImportStats()
    services_touches: set[str] = set()
    dates: list[datetime] = []

    # 1. Parser le fichier
    rows = parse_fis_sheet(file_bytes)
    stats.total_lignes = len(rows)

    if not rows:
        return FisImportResult(stats=stats)

    # 2. Pré-charger le cache des référentiels
    nature_cache: dict[str, Nature | None] = {}
    cause_cache: dict[str, tuple[Cause | None, bool]] = {}
    siege_cache: dict[str, Siege | None] = {}
    service_cache: dict[str, Service | None] = {}
    feeder_cache: dict[tuple[str, int | None], Feeder | None] = {}
    poste_source_cache: dict[str, PosteSource | None] = {}

    # 3. Importer ligne par ligne
    for row in rows:
        try:
            # Doublon
            if skip_existing and await _interruption_exists(db, row.numero):
                stats.doublons_ignores += 1
                continue

            # Résolution service
            if row.service_libelle not in service_cache:
                service_cache[row.service_libelle] = await _find_service(db, row.service_libelle)
            service = service_cache[row.service_libelle]
            service_id = service.id if service else None

            # Résolution poste source
            ps_key = row.source_libelle
            if ps_key not in poste_source_cache:
                poste_source_cache[ps_key] = await _find_poste_source(db, row.source_libelle)
            poste_source = poste_source_cache[ps_key]

            # Résolution feeder
            fdr_key = (row.feeder_libelle, service_id)
            if fdr_key not in feeder_cache:
                feeder_cache[fdr_key] = await _find_feeder(db, row.feeder_libelle, service_id)
            feeder = feeder_cache[fdr_key]

            # Résolution nature
            nat_key = _normalise(row.nature_libelle)
            if nat_key not in nature_cache:
                nature_cache[nat_key] = await _get_or_create_nature(db, row.nature_libelle)
            nature = nature_cache[nat_key]

            # Résolution cause
            cau_key = _normalise(row.cause_libelle)
            if cau_key not in cause_cache:
                cause_cache[cau_key] = await _get_or_create_cause(db, row.cause_libelle)
            cause, is_force_majeure = cause_cache[cau_key]

            # Résolution siège
            sie_key = _normalise(row.siege_libelle or "")
            if sie_key not in siege_cache:
                siege_cache[sie_key] = await _get_or_create_siege(db, row.siege_libelle)
            siege = siege_cache[sie_key]

            # Flags
            nature_code = nature.code if nature else ""
            is_fugitive = nature_code in NATURES_FUGITIVES
            is_programme = nature_code in NATURES_PROGRAMMEES

            # Conversion kW → MW, kWh → MWh
            puissance_mw = row.puissance_kw / 1000.0
            end_mwh = row.end_kwh / 1000.0

            # Créer l'interruption principale
            interruption = Interruption(
                numero_interruption=row.numero,
                type_direction="distribution",
                statut="valide",
                service_id=service_id,
                feeder_id=feeder.id if feeder else None,
                poste_source_id=poste_source.id if poste_source else None,
                date_apparition=row.date_apparition,
                date_remise=row.date_remise,
                duree_minutes=row.duree_minutes if row.duree_minutes > 0 else None,
                cause_id=cause.id if cause else None,
                siege_id=siege.id if siege else None,
                nature_id=nature.id if nature else None,
                is_force_majeure=is_force_majeure,
                is_fugitive=is_fugitive,
                puissance_coupee_mw=puissance_mw if puissance_mw > 0 else None,
                localisation=row.troncon,
                observations=row.observations,
                created_by_id=created_by_id,
            )
            db.add(interruption)
            await db.flush()  # obtenir l'id

            # Créer le détail distribution
            distrib = InterruptionDistribution(
                interruption_id=interruption.id,
                tension_kv=row.tension_kv,
                end_mwh=end_mwh if end_mwh > 0 else None,
                is_programme=is_programme,
            )
            db.add(distrib)
            await db.flush()

            # Statistiques
            stats.importees += 1
            stats.end_total_kwh += row.end_kwh
            dates.append(row.date_apparition)
            if row.service_libelle:
                services_touches.add(row.service_libelle)

        except Exception as exc:
            msg = f"Ligne N°{row.numero}: {exc}"
            logger.warning("FIS import error — %s", msg)
            stats.erreurs += 1
            stats.erreurs_detail.append(msg)
            continue

    return FisImportResult(
        stats=stats,
        periode_debut=min(dates) if dates else None,
        periode_fin=max(dates) if dates else None,
        services_touches=sorted(services_touches),
    )
