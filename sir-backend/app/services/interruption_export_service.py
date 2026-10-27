"""Export Excel / PDF de la liste des interruptions distribution (filtres de l'écran)."""

from __future__ import annotations

import io
from datetime import datetime

from app.schemas.interruption_distribution import InterruptionDistributionResponse

# Plafond d'une extraction synchrone : au-delà, filtrer davantage (période, région…)
MAX_LIGNES_EXPORT = 5000

_STATUTS = {"en_cours": "En cours", "cloture": "Clôturée", "valide": "Validée", "archive": "Archivée"}

COLONNES = [
    "N°",
    "Statut",
    "Apparition",
    "Remise",
    "Durée (min)",
    "Région",
    "Service",
    "Poste source",
    "Feeder",
    "Cause",
    "Nature",
    "END (MWh)",
    "Puissance coupée (MW)",
    "Clients MT",
    "Clients BT",
]


def _date(value: datetime | None) -> str:
    return value.strftime("%d/%m/%Y %H:%M") if value else ""


def ligne(item: InterruptionDistributionResponse) -> list[object]:
    return [
        item.numero_interruption,
        _STATUTS.get(item.statut, item.statut),
        _date(item.date_apparition),
        _date(item.date_remise),
        item.duree_minutes,
        item.region_libelle or "",
        item.service_libelle or "",
        item.poste_source_libelle or "",
        item.feeder_libelle or "",
        item.cause_libelle or "",
        item.nature_libelle or "",
        item.end_mwh,
        item.puissance_coupee_mw,
        item.nombre_clients_mt,
        item.nombre_clients_bt,
    ]


def construire_xlsx(items: list[InterruptionDistributionResponse], filtres: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Interruptions"
    ws.append([f"SIR — Interruptions distribution · {filtres}"])
    ws["A1"].font = Font(bold=True, size=12)
    ws.append([f"Extraction du {datetime.now():%d/%m/%Y %H:%M} · {len(items)} ligne(s)"])
    ws.append([])
    ws.append(COLONNES)
    for cell in ws[4]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="3C1053")
    for item in items:
        ws.append(ligne(item))
    largeurs = [10, 11, 17, 17, 11, 16, 26, 22, 20, 26, 16, 11, 12, 10, 10]
    for i, largeur in enumerate(largeurs, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largeur
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(COLONNES))}{4 + len(items)}"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def construire_pdf(items: list[InterruptionDistributionResponse], filtres: str) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4), leftMargin=24, rightMargin=24, topMargin=24, bottomMargin=24
    )
    styles = getSampleStyleSheet()
    # Colonnes utiles à la lecture papier (le détail complet est dans l'export Excel)
    entetes = ["N°", "Statut", "Apparition", "Remise", "Durée", "Service", "Feeder", "Cause", "END (MWh)"]
    data = [entetes]
    for item in items:
        data.append(
            [
                str(item.numero_interruption),
                _STATUTS.get(item.statut, item.statut),
                _date(item.date_apparition),
                _date(item.date_remise),
                f"{item.duree_minutes} min" if item.duree_minutes is not None else "",
                (item.service_libelle or "")[:28],
                (item.feeder_libelle or "")[:22],
                (item.cause_libelle or "")[:28],
                f"{item.end_mwh:.2f}".replace(".", ",") if item.end_mwh is not None else "",
            ]
        )
    table = Table(data, repeatRows=1, colWidths=[52, 58, 80, 80, 50, 140, 110, 150, 60])
    table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#3C1053")),
                ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                ("FONTSIZE", (0, 0), (-1, -1), 7.5),
                ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F3EEF8")]),
                ("ALIGN", (8, 1), (8, -1), "RIGHT"),
                ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#D4C8DE")),
            ]
        )
    )
    doc.build(
        [
            Paragraph("SIR — Interruptions distribution", styles["Title"]),
            Paragraph(f"{filtres} · {len(items)} ligne(s) · extraction du {datetime.now():%d/%m/%Y %H:%M}", styles["Normal"]),
            Spacer(1, 8),
            table,
        ]
    )
    return buffer.getvalue()
